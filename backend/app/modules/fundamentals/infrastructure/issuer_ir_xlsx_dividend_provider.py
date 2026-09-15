"""Issuer IR XLSX dividend provider (V2) — bounded RESEARCH / PARTIAL production.

Catalog (explicit adapters, not a universal crawler):

- Magnit (MGNT): dates + history XLSX join (board / approval / record / amount)
- Lukoil (LKOH): single declared-dividends XLSX (amount + record; optional AGM date)

``known_at`` quality is tracked explicitly:

- Magnit approval/board → ``APPROXIMATE_PUBLICATION_PROXY``
- Lukoil AGM decision date in period text → ``MEETING_DATE_PROXY``
- Lukoil record date only → ``RECORD_DATE_PROXY`` (conservative late PIT)

Status:

- shareholder approval / declared amount+record → ``APPROVED``
- only board recommendation → ``RECOMMENDED`` (not TR entitlement)

Never invents zero amounts. Never copies common→preferred. Skips incomplete rows.
"""

from __future__ import annotations

import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Literal

import httpx
from openpyxl import load_workbook

from app.modules.fundamentals.domain.types import (
    KNOWN_AT_QUALITY_APPROXIMATE_PUBLICATION_PROXY,
    SOURCE_ISSUER_IR_XLS_V1,
    DividendEventRef,
    DividendStatus,
)

MAGNIT_DATES_URL = (
    "https://www.magnit.com/files/ru/shareholders-and-investors/dividends-dates-0107-rus.xlsx"
)
MAGNIT_HISTORY_URL = (
    "https://www.magnit.com/files/ru/shareholders-and-investors/dividends-history-0107-rus.xlsx"
)
LUKOIL_DECLARED_URL = "https://lukoil.ru/FileSystem/9/729054.xlsx"

DEFAULT_TIMEOUT_S = 30.0

WorkbookFormat = Literal["magnit_dates_history", "lukoil_declared_single"]

_MONTHS_RU: dict[str, int] = {
    "января": 1,
    "февраля": 2,
    "марта": 3,
    "апреля": 4,
    "мая": 5,
    "июня": 6,
    "июля": 7,
    "августа": 8,
    "сентября": 9,
    "октября": 10,
    "ноября": 11,
    "декабря": 12,
}

_DATE_RE = re.compile(
    r"(?P<day>\d{1,2})\s+(?P<month>[А-Яа-яЁё]+)\s+(?P<year>\d{4})",
    re.UNICODE,
)
_AGM_DOT_DATE_RE = re.compile(r"от\s+(?P<day>\d{2})\.(?P<month>\d{2})\.(?P<year>\d{4})")
_AMOUNT_SHARE_RE = re.compile(
    r"(?P<amount>[\d\s]+(?:[.,]\d+)?)\s*(?P<cls>ао|ап|ao|ap)?",
    re.IGNORECASE | re.UNICODE,
)


@dataclass(frozen=True, slots=True)
class IssuerIrXlsxSpec:
    """One issuer's public IR workbook endpoints."""

    secid: str
    dates_url: str
    history_url: str
    workbook_format: WorkbookFormat = "magnit_dates_history"
    share_class: str = "common"  # common | preferred — never auto-copy across classes
    source_note: str = ""


DEFAULT_ISSUER_IR_CATALOG: tuple[IssuerIrXlsxSpec, ...] = (
    IssuerIrXlsxSpec(
        secid="MGNT",
        dates_url=MAGNIT_DATES_URL,
        history_url=MAGNIT_HISTORY_URL,
        workbook_format="magnit_dates_history",
        share_class="common",
        source_note="Magnit IR dates+history XLSX",
    ),
    IssuerIrXlsxSpec(
        secid="LKOH",
        dates_url=LUKOIL_DECLARED_URL,
        history_url=LUKOIL_DECLARED_URL,
        workbook_format="lukoil_declared_single",
        share_class="common",
        source_note="Lukoil IR declared dividends single XLSX (ао only)",
    ),
)


@dataclass(frozen=True, slots=True)
class PeriodKey:
    kind: str  # FY | M9 | H1 | Q1
    year: int

    def label(self) -> str:
        return f"{self.kind}:{self.year}"


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    return re.sub(r"\s+", " ", text)


def parse_russian_date(value: Any) -> date | None:
    """Parse Excel date / Russian textual date. Returns None when unparseable."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _clean_text(value)
    if not text:
        return None
    match = _DATE_RE.search(text)
    if not match:
        return None
    month = _MONTHS_RU.get(match.group("month").lower())
    if month is None:
        return None
    try:
        return date(int(match.group("year")), month, int(match.group("day")))
    except ValueError:
        return None


def normalize_period_label(raw: Any) -> PeriodKey | None:
    """Map Magnit dates/history period labels onto a join key.

    Dates sheet examples: ``2023 за год``, ``2020 9М``, ``2017 1П``, ``2012 1Кв``.
    History sheet examples: ``2023``, ``9М 2020``, ``1П 2017``, ``2020 ИТОГ`` (aggregate — skipped).
    """
    text = _clean_text(raw).lower().replace("ё", "е")
    if not text or "итог" in text:
        return None
    year_match = re.search(r"(20\d{2}|19\d{2})", text)
    if year_match is None:
        return None
    year = int(year_match.group(1))
    if re.search(r"9\s*м|9м", text):
        return PeriodKey("M9", year)
    if re.search(r"1\s*п|1п", text):
        return PeriodKey("H1", year)
    if re.search(r"1\s*кв|1кв", text):
        return PeriodKey("Q1", year)
    # ``за год`` / bare year → fiscal-year / year-end slice (not ИТОГ aggregates).
    return PeriodKey("FY", year)


def _sheet_rows(payload: bytes) -> list[tuple[Any, ...]]:
    wb = load_workbook(BytesIO(payload), data_only=True, read_only=True)
    try:
        ws = wb.active
        return [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_dates_sheet(payload: bytes) -> list[dict[str, Any]]:
    rows = _sheet_rows(payload)
    out: list[dict[str, Any]] = []
    for i, row in enumerate(rows):
        if i == 0 or not row:
            continue
        period_raw = row[1] if len(row) > 1 else None
        key = normalize_period_label(period_raw)
        if key is None:
            continue
        board = parse_russian_date(row[2] if len(row) > 2 else None)
        approval = parse_russian_date(row[3] if len(row) > 3 else None)
        record = parse_russian_date(row[4] if len(row) > 4 else None)
        out.append(
            {
                "period_raw": _clean_text(period_raw),
                "period_key": key,
                "board_recommendation_date": board,
                "shareholder_approval_date": approval,
                "record_date": record,
            }
        )
    return out


def parse_history_sheet(payload: bytes) -> dict[PeriodKey, float]:
    rows = _sheet_rows(payload)
    amounts: dict[PeriodKey, float] = {}
    for i, row in enumerate(rows):
        if i == 0 or not row:
            continue
        key = normalize_period_label(row[0] if len(row) > 0 else None)
        if key is None:
            continue
        raw_amount = row[3] if len(row) > 3 else None
        if raw_amount is None or raw_amount == "":
            continue
        try:
            amount = float(raw_amount)
        except (TypeError, ValueError):
            continue
        # Never invent zeros; skip non-positive as missing economics.
        if amount <= 0:
            continue
        amounts[key] = amount
    return amounts


def join_dividend_rows(
    dates_rows: Sequence[Mapping[str, Any]],
    history_amounts: Mapping[PeriodKey, float],
) -> list[dict[str, Any]]:
    joined: list[dict[str, Any]] = []
    for row in dates_rows:
        key: PeriodKey = row["period_key"]
        amount = history_amounts.get(key)
        joined.append({**dict(row), "amount_per_share": amount})
    return joined


def rows_to_event_refs(
    joined: Sequence[Mapping[str, Any]],
    *,
    issuer_id: int | None,
    instrument_id: int | None,
    secid: str,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[DividendEventRef]:
    events: list[DividendEventRef] = []
    for row in joined:
        board = row.get("board_recommendation_date")
        approval = row.get("shareholder_approval_date")
        record = row.get("record_date")
        amount = row.get("amount_per_share")
        if approval is not None:
            status = DividendStatus.APPROVED
            known_at = approval
        elif board is not None:
            status = DividendStatus.RECOMMENDED
            known_at = board
        else:
            continue

        if status == DividendStatus.APPROVED and (amount is None or record is None):
            # Do not invent entitlement economics.
            continue
        if status == DividendStatus.RECOMMENDED and amount is None and record is None:
            continue

        anchor = record or approval or board
        if date_from is not None and anchor is not None and anchor < date_from:
            continue
        if date_to is not None and anchor is not None and anchor > date_to:
            continue

        period_key: PeriodKey = row["period_key"]
        meta = {
            "known_at_quality": KNOWN_AT_QUALITY_APPROXIMATE_PUBLICATION_PROXY,
            "known_at_basis": (
                "shareholder_approval_date"
                if approval is not None
                else "board_recommendation_date"
            ),
            "known_at_note": (
                "Meeting/registry dates from issuer IR XLSX are not disclosure "
                "publication timestamps; used only as an approximate availability proxy."
            ),
            "secid": secid,
            "period_raw": row.get("period_raw"),
            "period_key": period_key.label(),
            "workbook_format": "magnit_dates_history",
            "share_class": "common",
            "entitlement_eligible": status == DividendStatus.APPROVED,
        }
        events.append(
            DividendEventRef(
                known_at=known_at,
                status=status,
                source=SOURCE_ISSUER_IR_XLS_V1,
                issuer_id=issuer_id,
                instrument_id=instrument_id,
                board_recommendation_date=board,
                shareholder_approval_date=approval,
                record_date=record,
                amount_per_share=float(amount) if amount is not None else None,
                currency="RUB",
                metadata=meta,
            )
        )
    return events


@dataclass
class IssuerIrXlsxDividendProvider:
    """Ports-compatible dividend provider for bounded IR XLSX catalog (MGNT+LKOH)."""

    source: str = SOURCE_ISSUER_IR_XLS_V1
    catalog: Sequence[IssuerIrXlsxSpec] = field(default_factory=lambda: DEFAULT_ISSUER_IR_CATALOG)
    issuer_id_by_secid: Mapping[str, int] = field(default_factory=dict)
    instrument_id_by_secid: Mapping[str, int] = field(default_factory=dict)
    # Optional local overrides for tests: secid → (dates_bytes_or_path, history_bytes_or_path)
    local_files: Mapping[str, tuple[Any, Any]] = field(default_factory=dict)
    timeout_s: float = DEFAULT_TIMEOUT_S
    trust_env_proxy: bool = False

    def configured_secids(self) -> frozenset[str]:
        return frozenset(spec.secid.upper() for spec in self.catalog)

    def _spec_for_secid(self, secid: str) -> IssuerIrXlsxSpec | None:
        needle = secid.upper()
        for spec in self.catalog:
            if spec.secid.upper() == needle:
                return spec
        return None

    def _spec_for_issuer(self, issuer_id: int) -> IssuerIrXlsxSpec | None:
        for secid, iid in self.issuer_id_by_secid.items():
            if int(iid) == int(issuer_id):
                return self._spec_for_secid(secid)
        return None

    def _download(self, url: str) -> bytes:
        with httpx.Client(
            timeout=self.timeout_s,
            follow_redirects=True,
            trust_env=self.trust_env_proxy,
        ) as client:
            response = client.get(url)
            response.raise_for_status()
            return response.content

    def _load_bytes(self, source: Any | None, *, url: str) -> bytes:
        if source is None:
            return self._download(url)
        if isinstance(source, bytes | bytearray):
            return bytes(source)
        if isinstance(source, Path):
            return source.read_bytes()
        if isinstance(source, str):
            if source.startswith("http://") or source.startswith("https://"):
                return self._download(source)
            return Path(source).read_bytes()
        raise TypeError(f"Unsupported workbook source type: {type(source)!r}")

    def _load_workbooks(self, spec: IssuerIrXlsxSpec) -> tuple[bytes, bytes]:
        local = self.local_files.get(spec.secid.upper()) or self.local_files.get(spec.secid)
        if local is not None:
            dates_src, history_src = local
            return (
                self._load_bytes(dates_src, url=spec.dates_url),
                self._load_bytes(history_src, url=spec.history_url),
            )
        return (
            self._load_bytes(None, url=spec.dates_url),
            self._load_bytes(None, url=spec.history_url),
        )

    def parse_secid_events(
        self,
        secid: str,
        *,
        dates_payload: bytes,
        history_payload: bytes,
        date_from: date | None = None,
        date_to: date | None = None,
        workbook_format: WorkbookFormat | None = None,
        share_class: str = "common",
    ) -> list[DividendEventRef]:
        fmt = workbook_format or "magnit_dates_history"
        if fmt == "lukoil_declared_single":
            from app.modules.fundamentals.infrastructure.lukoil_ir_xlsx import (
                lukoil_rows_to_event_refs,
                parse_lukoil_declared_sheet,
            )

            rows = parse_lukoil_declared_sheet(
                dates_payload, expected_share_class=share_class
            )
            return lukoil_rows_to_event_refs(
                rows,
                issuer_id=self.issuer_id_by_secid.get(secid.upper()),
                instrument_id=self.instrument_id_by_secid.get(secid.upper()),
                secid=secid.upper(),
                date_from=date_from,
                date_to=date_to,
            )
        dates_rows = parse_dates_sheet(dates_payload)
        history = parse_history_sheet(history_payload)
        joined = join_dividend_rows(dates_rows, history)
        return rows_to_event_refs(
            joined,
            issuer_id=self.issuer_id_by_secid.get(secid.upper()),
            instrument_id=self.instrument_id_by_secid.get(secid.upper()),
            secid=secid.upper(),
            date_from=date_from,
            date_to=date_to,
        )

    def fetch_by_secid(
        self,
        secid: str,
        *,
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> Sequence[DividendEventRef]:
        spec = self._spec_for_secid(secid)
        if spec is None:
            return ()
        dates_payload, history_payload = self._load_workbooks(spec)
        return self.parse_secid_events(
            spec.secid,
            dates_payload=dates_payload,
            history_payload=history_payload,
            date_from=date_from,
            date_to=date_to,
            workbook_format=spec.workbook_format,
            share_class=spec.share_class,
        )

    def fetch_dividends(
        self,
        issuer: int,
        *,
        date_from: date | None = None,
        date_to: date | None = None,
    ) -> Sequence[DividendEventRef]:
        spec = self._spec_for_issuer(issuer)
        if spec is None:
            for secid, iid in self.issuer_id_by_secid.items():
                if int(iid) == int(issuer):
                    return self.fetch_by_secid(secid, date_from=date_from, date_to=date_to)
            return ()
        return self.fetch_by_secid(spec.secid, date_from=date_from, date_to=date_to)

    def readiness(self) -> dict[str, Any]:
        return {
            "status": "PARTIAL",
            "provider": self.source,
            "accepted": True,
            "verdict": "PARTIAL_RESEARCH_PRODUCTION_BOUNDED",
            "bounded_secids": sorted(self.configured_secids()),
            "catalog": [
                {
                    "secid": s.secid,
                    "format": s.workbook_format,
                    "share_class": s.share_class,
                    "note": s.source_note,
                    "dates_url": s.dates_url,
                    "history_url": s.history_url,
                }
                for s in self.catalog
            ],
            "reasons": [
                "moex_iss_dividends_rejected",
                "e_disclosure_403_rejected",
                "issuer_ir_xlsx_bounded_multi_issuer",
                "known_at_not_exact_publication_timestamp",
            ],
            "notes": [
                "MGNT: dates+history join; known_at=APPROXIMATE_PUBLICATION_PROXY.",
                "LKOH: declared sheet; known_at=MEETING_DATE_PROXY or RECORD_DATE_PROXY.",
                "RECOMMENDED (board-only) is not TR entitlement.",
                "Common/preferred never auto-copied.",
            ],
            "as_of": date.today().isoformat(),
        }
